"""Acrescenta ao briefing dos mesários os procedimentos de encerramento.

Uso: python3 build.py ORIGEM.pptx DESTINO.pptx
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

NAVY, GREEN, RED, GOLD = "042B5A", "5F882E", "C8102E", "E8C63A"
BEIGE, PANEL, LINE, GRAY, WHITE, BLUE = "F2EFE2", "EDEADB", "E4DFCC", "6B6B6B", "FFFFFF", "33507E"
FONT = "Montserrat"


def rgb(h):
    return RGBColor.from_string(h)


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, adj=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    if adj is not None:
        s.adjustments[0] = adj
    return s


def text(slide, x, y, w, h, runs, size=10, color=GRAY, bold=False, align=None,
         anchor=MSO_ANCHOR.TOP, spacing=None):
    """runs: str, or list of paragraphs; each paragraph is a str or list of (txt, bold, color)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        segs = [(para, bold, color)] if isinstance(para, str) else para
        for txt, b, c in segs:
            r = p.add_run()
            r.text = txt
            f = r.font
            f.name, f.size, f.bold = FONT, Pt(size), b
            f.color.rgb = rgb(c)
    return tb


def badge(slide, x, y, d, label, fill, size=10):
    box(slide, x, y, d, d, fill, MSO_SHAPE.OVAL)
    text(slide, x, y, d, d, label, size=size, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def header(slide, kicker, title, sub):
    box(slide, 0, 0, 13.333, 7.5, BEIGE)
    box(slide, 0, 0, 13.333, 0.13, NAVY)
    text(slide, 0.62, 0.42, 8.0, 0.2, kicker, size=9, color=GREEN, bold=True)
    text(slide, 0.62, 0.66, 11.5, 0.56, title, size=30, color=NAVY, bold=True)
    text(slide, 0.62, 1.26, 11.6, 0.42, sub, size=12, color=GRAY)


def footnote(slide, y, msg):
    text(slide, 0.62, y, 12.09, 0.4, msg, size=7.5, color=GRAY)


# --------------------------------------------------------------------------
# Slide A · encerramento da votação, passo a passo
# --------------------------------------------------------------------------
def slide_encerramento(prs, layout):
    s = prs.slides.add_slide(layout)
    header(s, "ENCERRAMENTO",
           "Cinco passos para fechar a seção",
           "A votação só pode ser encerrada a partir das 17h, no horário de Dublin, e com a fila vazia. "
           "Quem já estava na fila às 17h vota.")

    passos = [
        ("1", GREEN, "Senhas", "SECRETÁRIO",
         "Às 17h, entrega senhas a quem está na fila, da última pessoa até a primeira. "
         "Quem chegar depois não entra."),
        ("2", NAVY, "Encerrar a votação", "PRESIDENTE",
         "No Terminal do Mesário: CORRIGE › outras opções › 2 – Encerrar votação. "
         "Digita o próprio título, CONFIRMA e registra de novo os mesários presentes."),
        ("3", NAVY, "Boletins", "A URNA IMPRIME",
         "Confere a 1ª via do BU e, se legível, CONFIRMA. Saem a 2ª à 5ª via do BU, "
         "o BUJ e o BIM. Até 5 vias extras do BU, se pedidas."),
        ("4", GREEN, "Assinaturas", "TODA A MESA",
         "BU e BUJ: todas as vias, pela mesa e pelos fiscais presentes. "
         "BIM: só pela mesa."),
        ("5", RED, "Mídia de Resultado", "PRESIDENTE",
         "Só retira a MR depois da mensagem na tela. Entrega à equipe de transmissão, "
         "recoloca o lacre e desliga a urna."),
    ]
    gap, x0, w = 0.18, 0.62, (12.09 - 4 * 0.18) / 5
    for i, (n, cor, tit, dono, desc) in enumerate(passos):
        x = x0 + i * (w + gap)
        box(s, x, 1.86, w, 2.62, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, LINE, 0.03)
        badge(s, x + 0.24, 2.08, 0.34, n, cor)
        text(s, x + 0.24, 2.54, w - 0.44, 0.34, tit, size=12.5, color=NAVY, bold=True)
        text(s, x + 0.24, 2.9, w - 0.44, 0.22, dono, size=7, color=GREEN, bold=True)
        text(s, x + 0.24, 3.16, w - 0.44, 1.26, desc, size=8.5, color=GRAY)

    # destino das cinco vias obrigatórias do BU
    box(s, 0.62, 4.72, 7.3, 1.98, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, LINE, 0.03)
    text(s, 0.86, 4.88, 6.9, 0.24, "AS 5 VIAS OBRIGATÓRIAS DO BOLETIM DE URNA",
         size=8, color=GREEN, bold=True)
    vias = [
        ("2 vias", "no envelope, entregues à Embaixada para a Justiça Eleitoral"),
        ("1 via", "afixada em local visível, na entrada da seção"),
        ("1 via", "com o Presidente, para conferir o resultado no portal do TSE"),
        ("1 via", "às duas últimas pessoas a votar; as extras, a fiscais, MP e imprensa, se pedirem"),
    ]
    for i, (q, d) in enumerate(vias):
        y = 5.2 + i * 0.36
        text(s, 0.86, y, 0.8, 0.3, q, size=10.5, color=NAVY, bold=True)
        text(s, 1.72, y + 0.02, 6.05, 0.3, d, size=9.5, color=GRAY)

    box(s, 8.2, 4.72, 4.51, 1.98, NAVY, MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.03)
    text(s, 8.48, 4.92, 4.0, 0.3, "Porta aberta, duas testemunhas", size=14, color=GOLD, bold=True)
    text(s, 8.48, 5.32, 4.0, 1.3,
         "O encerramento é público. As duas últimas pessoas a votar são convidadas a "
         "acompanhar e recebem uma via do BU. A participação ou a recusa vai para a ata.",
         size=10.5, color=WHITE)

    footnote(s, 6.86,
             "Manual do Mesário TSE 2026, p. 24–26; Manual do Representante Diplomático TSE 2026, p. 7 e 18–20. "
             "O Manual do Mesário fala em 17h de Brasília; no exterior vale o horário local. O BEHB citado no manual "
             "nacional não se aplica: no exterior não há identificação biométrica.")
    return s


# --------------------------------------------------------------------------
# Slide B · conferência e liberação
# --------------------------------------------------------------------------
def slide_liberacao(prs, layout):
    s = prs.slides.add_slide(layout)
    header(s, "ANTES DE IR EMBORA",
           "Ninguém sai antes da conferência",
           "A mesa só é liberada depois que a equipe da Embaixada confere o fechamento da seção "
           "e recebe, em mãos, estes quatro documentos.")

    docs = [
        ("Ata", GREEN, "Ata da Mesa", "VIA ÚNICA",
         "Completa, com todas as ocorrências do dia. Assinada pelos mesários e pelos fiscais presentes."),
        ("BIM", NAVY, "Boletim de Identificação de Mesários", "VIA ÚNICA · IMPRESSO PELA URNA",
         "Traz os mesários registrados no terminal. Assinado só pela mesa."),
        ("BU", NAVY, "Boletins de Urna", "AS 2 VIAS DO ENVELOPE",
         "As vias destinadas à Justiça Eleitoral. Assinadas pela mesa e pelos fiscais presentes."),
        ("BUJ", RED, "Boletim de Justificativa", "VIA ÚNICA · IMPRESSO PELA URNA",
         "Lista as justificativas registradas na urna. Assinado pela mesa e pelos fiscais presentes."),
    ]
    gap, w = 0.3, (12.09 - 3 * 0.3) / 4
    for i, (sig, cor, tit, sub, desc) in enumerate(docs):
        x = 0.62 + i * (w + gap)
        box(s, x, 1.86, w, 2.5, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, LINE, 0.03)
        badge(s, x + 0.3, 2.08, 0.6, sig, cor, size=11 if len(sig) <= 2 else 9.5)
        text(s, x + 0.3, 2.84, w - 0.56, 0.56, tit, size=13, color=NAVY, bold=True,
             anchor=MSO_ANCHOR.BOTTOM)
        text(s, x + 0.3, 3.46, w - 0.56, 0.22, sub, size=7, color=GREEN, bold=True)
        text(s, x + 0.3, 3.72, w - 0.56, 0.6, desc, size=9, color=GRAY)

    # o que a equipe da Embaixada confere, além dos papéis
    text(s, 0.62, 4.56, 8.0, 0.22, "A CONFERÊNCIA NA MESA", size=8, color=GREEN, bold=True)
    checks = [
        "Todas as vias assinadas, por quem deve assinar",
        "Uma via do BU afixada na entrada da seção",
        "Mídia de Resultado entregue à transmissão",
        "Urna desligada e guardada na caixa",
        "“NC” no caderno, no lugar da assinatura de quem faltou",
        "Comprovantes destacados por engano junto do material",
    ]
    for i, c in enumerate(checks):
        col, row = i % 2, i // 2
        x, y = 0.62 + col * 4.05, 4.88 + row * 0.42
        box(s, x, y + 0.02, 0.26, 0.26, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, GREEN, 0.15)
        text(s, x + 0.4, y + 0.04, 3.55, 0.3, c, size=9.5, color=NAVY)

    box(s, 8.9, 4.56, 3.81, 1.74, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.03)
    box(s, 8.9, 4.56, 0.09, 1.74, RED)
    text(s, 9.18, 4.72, 3.4, 0.3, "Por que esperar?", size=12.5, color=NAVY, bold=True)
    text(s, 9.18, 5.06, 3.4, 1.2,
         "Esses papéis são a prova do resultado da seção. Uma assinatura que falte "
         "não se colhe depois que a mesa foi embora.",
         size=9.5, color=NAVY)

    footnote(s, 6.5,
             "Manual do Mesário TSE 2026, p. 15 e 26–27 (envio de BU, BUJ, BIM e Ata em envelope) e p. 32 "
             "(modelo da ata: saída antecipada é a ocorrência 3.2). Manual do Representante Diplomático TSE 2026, "
             "p. 19–20 e 25. A liberação só após a conferência é regra do Posto, não do manual.")
    return s


# --------------------------------------------------------------------------
# Agenda (slide 2): seis cartões em 3×2 viram sete em 4+3
# --------------------------------------------------------------------------
def atualiza_agenda(slide):
    shapes = list(slide.shapes)
    sub = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.startswith("Seis assuntos"))
    r = sub.text_frame.paragraphs[0].runs
    r[0].text = ("Sete assuntos. Os quatro primeiros são sobre você; os dois seguintes, "
                 "sobre o eleitor; o último, sobre o fim do dia.")
    for extra in r[1:]:
        extra.text = ""
    sub.width = Inches(12.09)

    # cada cartão: fundo, círculo, número, título, descrição (5 formas seguidas)
    first = shapes.index(next(sh for sh in shapes if sh.name == "Shape 5"))
    cards = [shapes[first + 5 * k: first + 5 * k + 5] for k in range(6)]

    # sétimo cartão: cópia do sexto
    tree = slide.shapes._spTree
    novo = []
    for sh in cards[5]:
        el = copy.deepcopy(sh._element)
        tree.append(el)
        novo.append(el)
    next_id = max(int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))) + 1
    for el in novo:
        el.find(".//" + qn("p:cNvPr")).set("id", str(next_id))
        next_id += 1
    cards.append([sh for sh in slide.shapes if sh._element in novo])

    def set_text(sh, t):
        runs = sh.text_frame.paragraphs[0].runs
        runs[0].text = t
        for extra in runs[1:]:
            extra.text = ""

    set_text(cards[6][2], "7")
    set_text(cards[6][3], "Encerramento")
    set_text(cards[6][4], "Como a mesa fecha a seção — e os quatro documentos que liberam a equipe.")

    gap, w = 0.3, (12.09 - 3 * 0.3) / 4
    for k, card in enumerate(cards):
        row, col = (0, k) if k < 4 else (1, k - 4)
        x, y = 0.62 + col * (w + gap), 2.3 + row * 2.48
        bg, dot, num, tit, desc = card
        bg.left, bg.top, bg.width = Inches(x), Inches(y), Inches(w)
        for sh in (dot, num):
            sh.left, sh.top = Inches(x + 0.34), Inches(y + 0.34)
        tit.left, tit.top, tit.width = Inches(x + 0.34), Inches(y + 0.86), Inches(w - 0.68)
        desc.left, desc.top = Inches(x + 0.34), Inches(y + 1.24)
        desc.width, desc.height = Inches(w - 0.68), Inches(0.75)


def move_slide(prs, slide, new_index):
    lst = prs.slides._sldIdLst
    el = next(e for e in lst if prs.slides.part.related_part(e.rId) is slide.part)
    lst.remove(el)
    lst.insert(new_index, el)


def main(src, dst):
    prs = Presentation(src)
    assert len(prs.slides) == 15, "esperava o briefing de 15 slides"
    layout = prs.slides[14].slide_layout
    atualiza_agenda(prs.slides[1])
    a = slide_encerramento(prs, layout)
    b = slide_liberacao(prs, layout)
    # entram depois de "As primeiras horas decidem o dia" (14) e antes do checklist final
    move_slide(prs, a, 14)
    move_slide(prs, b, 15)
    prs.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
