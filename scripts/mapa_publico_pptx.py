#!/usr/bin/env python3
"""O mapa simplificado em PPTX editável — a mesma folha, em quatro slides.

A folha `mapa/mapa_publico.html` é para a internet; esta é para a sala. Sai do
**mesmo dado** (`separadores_fila`, `prancheta_hall2.json`, `decisoes.json`,
`paleta.json`, `data/ring3_montagem.json`) e diz exatamente as mesmas coisas.

**Tudo é forma nativa** — retângulo, linha, caixa de texto. Nada de imagem
colada: quem receber o arquivo move uma mesa, troca uma cor ou corrige um
número sem voltar aqui. É o que "editável" tem de querer dizer.

    python3 scripts/mapa_publico_pptx.py            confere que dá para gerar
    python3 scripts/mapa_publico_pptx.py --grava    grava saidas/onde_voce_vota.pptx

A fonte pedida é a **Montserrat**, da campanha. Quem não a tiver instalada vê a
substituta do sistema; o arquivo não deixa de abrir por causa disso.
"""
import math
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt
from lxml import etree

import mapa_publico as MP
import separadores_fila as SF

SAIDA = MP.RAIZ / "saidas" / "onde_voce_vota.pptx"
W_SLIDE, H_SLIDE = 13.333, 7.5
FONTE = "Montserrat"
# caracteres por linha a 10,5 pt numa caixa de 6,3 pol — medido, nao estimado
# por cima: e o que decide quanto cada passo empurra o seguinte para baixo.
CHARS_LINHA = 76

C = lambda h: RGBColor.from_string(h.lstrip("#").upper())
TINTA, CREME, CINZA = C(MP.MARINHO), C(MP.CREME), C(MP.CINZA)
REGUA, VERDE, PREF = C(MP.REGUA), C(MP.VERDE), C(MP.PREF)
GRAFITE = C(MP.GRAFITE)
ZONA = {k: C(MP.ZONA[k]) for k in "ABC"}
TINTA_Z = {k: C(MP.TINTA[k]) for k in "ABC"}
# tinta de zona sobre fundo claro: o amarelo do rolo não se lê em texto
TXT_Z = {"A": C(MP.ZONA["A"]), "B": C("#9A7C00"), "C": C(MP.ZONA["C"])}
BRANCO = C("#FFFFFF")


# --------------------------------------------------------------------------
# Primitivas
# --------------------------------------------------------------------------
def fundo(slide, cor=CREME):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                               Inches(W_SLIDE), Inches(H_SLIDE))
    r.fill.solid(); r.fill.fore_color.rgb = cor
    r.line.fill.background(); r.shadow.inherit = False
    return r


def caixa(slide, x, y, w, h, texto, tam=12, cor=TINTA, negrito=False,
          alinha=PP_ALIGN.LEFT, esp=None, maiuscula=False, ancora=MSO_ANCHOR.TOP,
          giro=None, entrelinha=1.18):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, linha in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinha
        p.line_spacing = entrelinha
        r = p.add_run()
        r.text = linha.upper() if maiuscula else linha
        f = r.font
        f.name = FONTE; f.size = Pt(tam); f.bold = negrito; f.color.rgb = cor
        if esp is not None:                      # letter-spacing, em centésimos de ponto
            r.font._rPr.set("spc", str(int(esp * 100)))
    if giro:
        tf._bodyPr.set("rot", str(int(giro * 60000)))
    return tb


def retangulo(slide, x, y, w, h, cor=None, alpha=None, borda=None, larg=0.75,
              raio=None, tracejado=False):
    forma = MSO_SHAPE.ROUNDED_RECTANGLE if raio else MSO_SHAPE.RECTANGLE
    r = slide.shapes.add_shape(forma, Inches(x), Inches(y),
                               Inches(max(w, 0.004)), Inches(max(h, 0.004)))
    if raio:
        r.adjustments[0] = raio
    if cor is None:
        r.fill.background()
    else:
        r.fill.solid(); r.fill.fore_color.rgb = cor
        if alpha is not None:
            _transparencia(r.fill.fore_color, alpha)
    if borda is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = borda; r.line.width = Pt(larg)
        if tracejado:
            r.line._get_or_add_ln().append(
                _el('a:prstDash', {'val': 'dash'}))
    r.shadow.inherit = False
    return r


def linha(slide, x1, y1, x2, y2, cor=TINTA, larg=1.0, tracejado=False, seta=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
    ln.line.color.rgb = cor; ln.line.width = Pt(larg)
    el = ln.line._get_or_add_ln()
    if tracejado:
        el.append(_el('a:prstDash', {'val': 'dash'}))
    if seta:
        el.append(_el('a:tailEnd', {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    ln.shadow.inherit = False
    return ln


NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _el(tag, attrs):
    e = etree.SubElement(etree.Element("x"), f"{{{NS}}}{tag.split(':')[1]}")
    for k, v in attrs.items():
        e.set(k, v)
    return e


def _transparencia(cor, alpha):
    """alpha 0..1 = quanto some. O python-pptx não expõe, então vai no XML."""
    srgb = cor._xFill.find(f"{{{NS}}}srgbClr")
    srgb.append(_el('a:alpha', {'val': str(int((1 - alpha) * 100000))}))


def rotulo_secao(slide, x, y, texto, cor=VERDE):
    caixa(slide, x, y, 4.0, 0.2, texto, tam=8, cor=cor, negrito=True, esp=1.6,
          maiuscula=True)


# --------------------------------------------------------------------------
# Slide 1 — capa
# --------------------------------------------------------------------------
def capa(prs, dados):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s)
    retangulo(s, 0, 0, W_SLIDE, 0.13, cor=TINTA)
    rotulo_secao(s, 0.9, 1.15, "Eleições 2026 · 1º turno · domingo, 4 de outubro · 8h às 17h")
    caixa(s, 0.9, 1.5, 9.4, 1.6, "Onde você vota\nno RDS", tam=54, negrito=True,
          entrelinha=0.98)
    caixa(s, 0.9, 3.5, 7.4, 1.5,
          "Todas as seções da Irlanda votam num único lugar: o Hall 2 do RDS, na "
          "Merrion Road, em Ballsbridge. O caminho é o mesmo para todo mundo — muda "
          "só uma coisa, e é a letra da sua porta.", tam=15, cor=CINZA)
    x = 0.9
    for valor, rot in ((dados["secoes"], "seções"), (dados["urnas"], "urnas"),
                       (dados["aptos"], "eleitores aptos"), ("3", "portas: A, B e C")):
        caixa(s, x, 5.35, 2.4, 0.5, str(valor), tam=30, negrito=True)
        caixa(s, x, 5.92, 2.4, 0.3, rot, tam=10, cor=CINZA, negrito=True, esp=1.2,
              maiuscula=True)
        x += 2.55
    for i, k in enumerate("ABC"):
        px = 10.6 + i * 0.92
        retangulo(s, px, 1.5, 0.78, 0.78, cor=ZONA[k], raio=0.14)
        caixa(s, px, 1.68, 0.78, 0.5, k, tam=34, negrito=True, cor=TINTA_Z[k],
              alinha=PP_ALIGN.CENTER)
        caixa(s, px, 2.38, 0.78, 0.5, MP.PAREDE[k], tam=9, cor=CINZA,
              alinha=PP_ALIGN.CENTER)
    caixa(s, 10.6, 2.95, 2.6, 0.9,
          "Cada porta serve uma parede inteira, e só ela.", tam=11, cor=CINZA)
    caixa(s, 0.9, 6.85, 11.5, 0.3,
          "Royal Dublin Society, Hall 2 · Merrion Road, Ballsbridge, Dublin 4, D04 AK83"
          " · entrada de eleitores pelo portão B", tam=9, cor=CINZA)
    return s


# --------------------------------------------------------------------------
# Slide 2 — o caminho
# --------------------------------------------------------------------------
def slide_caminho(prs, planta, dec, dados):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s)
    rotulo_secao(s, 0.62, 0.42, "Passo a passo")
    caixa(s, 0.62, 0.68, 8.0, 0.5, "Da calçada até a urna", tam=27, negrito=True)

    dx, apron = MP.encaixe(planta, dec)
    RW, RD = MP.RING["ring"]
    COR = MP.RING["corredor"]
    PROF, UTIL = RD - COR, RW - COR
    LARG, ALT = planta["salao"]["largura"], planta["salao"]["altura"]
    rx = lambda x: x + dx
    ry = lambda y: -apron - y

    RUA_X, RUA_W, PORTAO_Y = 60.0, 8.0, 33.0
    CALCADA = RUA_X - 1.6
    x0, x1 = -3.0, RUA_X + RUA_W
    y0, y1 = ry(RD) - 4.0, ALT + 4.0
    LX, LY, LW, LH = 0.62, 1.32, 5.3, 5.75
    e = min(LW / (x1 - x0), LH / (y1 - y0))
    X = lambda x: LX + (x - x0) * e
    Y = lambda y: LY + (y1 - y) * e

    # rua e calçada
    retangulo(s, X(RUA_X), Y(y1), RUA_W * e, (y1 - y0) * e, cor=C("#D8D3C3"))
    caixa(s, X(RUA_X) - 0.45, Y(ALT - 6), 1.4, 0.2, "MERRION ROAD", tam=7.5,
          cor=CINZA, negrito=True, esp=1.4, alinha=PP_ALIGN.CENTER, giro=-90)
    linha(s, X(CALCADA), Y(y1 - 1), X(CALCADA), Y(y0 + 1), cor=REGUA, larg=2)

    # salão, apron, pátio
    retangulo(s, X(0), Y(ALT), LARG * e, ALT * e, cor=BRANCO, borda=TINTA, larg=1.4)
    caixa(s, X(0), Y(ALT / 2 + 1.5), LARG * e, 0.3, "HALL 2", tam=15, negrito=True,
          alinha=PP_ALIGN.CENTER)
    caixa(s, X(0), Y(ALT / 2 - 1.0), LARG * e, 0.24, "o salão de votação", tam=9,
          cor=CINZA, alinha=PP_ALIGN.CENTER)
    caixa(s, X(0), Y(ALT / 2 - 4.0), LARG * e, 0.24,
          "as 28 mesas ficam nas paredes — ver o slide seguinte", tam=7.5, cor=CINZA,
          alinha=PP_ALIGN.CENTER)
    retangulo(s, X(rx(0)), Y(0), RW * e, apron * e, cor=C("#E4E0CF"))
    caixa(s, X(rx(1.5)), Y(-apron / 2 + 1.2), 1.6, 0.2,
          f"apron · {MP.br1(apron)} m", tam=7.5, cor=CINZA)
    retangulo(s, X(rx(0)), Y(ry(0)), RW * e, RD * e, cor=C("#EDEADB"),
              borda=CINZA, larg=0.75, tracejado=True)
    retangulo(s, X(rx(UTIL)), Y(ry(0)), COR * e, RD * e, cor=C("#DBD6C2"))
    retangulo(s, X(rx(0)), Y(ry(PROF)), UTIL * e, COR * e, cor=C("#DBD6C2"))
    caixa(s, X(rx(0)), Y(ry(PROF + COR / 2) + 0.9), UTIL * e, 0.2,
          "o corredor distribui as três filas: C, depois B, depois A", tam=7,
          cor=CINZA, alinha=PP_ALIGN.CENTER)

    for k in "ABC":
        za = MP.RING["x0"][k]; zb = za + MP.RING["larguras"][k]
        retangulo(s, X(rx(za)), Y(ry(0)), (zb - za) * e, PROF * e, cor=ZONA[k],
                  alpha=0.78, borda=ZONA[k], larg=0.75)
        caixa(s, X(rx(za)), Y(ry(PROF / 2) + 1.6), (zb - za) * e, 0.4, k, tam=22,
              negrito=True, cor=TXT_Z[k], alinha=PP_ALIGN.CENTER)
        caixa(s, X(rx(za)), Y(ry(PROF / 2) - 1.4), (zb - za) * e, 0.2,
              f"fila da porta {k}", tam=7, cor=CINZA, alinha=PP_ALIGN.CENTER)
        ba, bb = MP.RING["bocas_fundo"][k]["x"]
        linha(s, X(rx(ba)), Y(ry(PROF)), X(rx(bb)), Y(ry(PROF)), cor=ZONA[k], larg=3)
        sa, sb = MP.RING["saidas_norte"][k]
        linha(s, X(rx(sa)), Y(ry(0)), X(rx(sb)), Y(ry(0)), cor=ZONA[k], larg=3)
        linha(s, X(rx((sa + sb) / 2)), Y(ry(0)), X(MP.RING["portas_eixo"][k] + dx),
              Y(0), cor=ZONA[k], larg=1.6, seta=True)

    for p in planta["portas"]:
        sin = dec["sinalizacao_portas"].get(p["id"], {})
        papel = sin.get("papel")
        if papel == "entrada":
            k = sin["entrada"]
            linha(s, X(p["x1"]), Y(0), X(p["x2"]), Y(0), cor=ZONA[k], larg=4.5)
            caixa(s, X((p["x1"] + p["x2"]) / 2) - 0.2, Y(0) - 0.24, 0.4, 0.22, k,
                  tam=12, negrito=True, cor=TXT_Z[k], alinha=PP_ALIGN.CENTER)
        elif papel == "preferencial":
            linha(s, X(p["x1"]), Y(0), X(p["x2"]), Y(0), cor=PREF, larg=4.5)
            caixa(s, X(p["x2"]) + 0.03, Y(0) - 0.22, 0.9, 0.2, "preferencial",
                  tam=7, cor=PREF, negrito=True)
        elif papel == "saida":
            linha(s, X(p["x1"]), Y(0), X(p["x2"]), Y(0), cor=GRAFITE, larg=4.5)
            caixa(s, X((p["x1"] + p["x2"]) / 2) - 0.25, Y(0) + 0.03, 0.5, 0.18,
                  "saída", tam=6.5, cor=CINZA, alinha=PP_ALIGN.CENTER)

    # a rota, do portão ao canto nordeste
    cx = LARG + 4.0
    rota = [(CALCADA, PORTAO_Y), (cx, PORTAO_Y), (cx, ry(0) + 1.5),
            (rx(UTIL + COR / 2), ry(0) + 1.5), (rx(UTIL + COR / 2), ry(3.0))]
    for a, b in zip(rota, rota[1:]):
        linha(s, X(a[0]), Y(a[1]), X(b[0]), Y(b[1]), cor=VERDE, larg=2.2,
              tracejado=True, seta=(b is rota[-1]))
    caixa(s, X(cx) - 0.62, Y(14), 1.4, 0.2, "lateral leste do Hall 2", tam=7,
          cor=CINZA, alinha=PP_ALIGN.CENTER, giro=-90)
    r = 0.19
    retangulo(s, X(CALCADA) - r / 2, Y(PORTAO_Y) - r / 2, r, r, cor=VERDE, raio=0.5)
    caixa(s, X(CALCADA) - 1.72, Y(PORTAO_Y) - 0.13, 1.5, 0.2, "PORTÃO B", tam=9,
          negrito=True, cor=VERDE, alinha=PP_ALIGN.RIGHT)
    caixa(s, X(CALCADA) - 1.72, Y(PORTAO_Y) + 0.05, 1.5, 0.2, "entrada do eleitor",
          tam=7, cor=CINZA, alinha=PP_ALIGN.RIGHT)

    zbc = rx(MP.RING["x0"]["B"] + MP.RING["larguras"]["B"] / 2)
    passos = [(CALCADA - 12.0, PORTAO_Y, "1"), (cx + 2.5, PORTAO_Y - 22, "2"),
              (rx(RW) + 2.6, ry(9), "3"), (zbc, ry(26.5), "4"),
              (MP.RING["portas_eixo"]["B"] + dx, -3.6, "5")]
    d = 0.26
    for px, py, n in passos:
        retangulo(s, X(px) - d / 2, Y(py) - d / 2, d, d, cor=VERDE, raio=0.5)
        caixa(s, X(px) - d / 2, Y(py) - 0.085, d, 0.2, n, tam=10, negrito=True,
              cor=BRANCO, alinha=PP_ALIGN.CENTER)

    # escala
    linha(s, X(x0 + 2), Y(y0 + 3), X(x0 + 22), Y(y0 + 3), cor=TINTA, larg=1.2)
    caixa(s, X(x0 + 22) + 0.05, Y(y0 + 3) - 0.08, 0.5, 0.18, "20 m", tam=7, cor=CINZA)

    # os cinco passos, à direita
    px, py = 6.3, 1.32
    for i, (titulo, corpo) in enumerate(dados["passos"]):
        retangulo(s, px, py + 0.02, 0.26, 0.26, cor=VERDE, raio=0.5)
        caixa(s, px, py + 0.055, 0.26, 0.2, str(i + 1), tam=10, negrito=True,
              cor=BRANCO, alinha=PP_ALIGN.CENTER)
        caixa(s, px + 0.42, py, 6.3, 0.26, titulo, tam=13, negrito=True)
        linhas = math.ceil(len(corpo) / CHARS_LINHA)
        caixa(s, px + 0.42, py + 0.3, 6.3, 0.2 * linhas, corpo, tam=10.5, cor=CINZA)
        py += 0.36 + 0.2 * linhas + 0.14
    retangulo(s, px, py + 0.12, 6.72, 0.62, cor=BRANCO, borda=REGUA)
    retangulo(s, px, py + 0.12, 0.05, 0.62, cor=PREF)
    caixa(s, px + 0.22, py + 0.26, 6.3, 0.4,
          "Atendimento preferencial: idoso, gestante, pessoa com deficiência e quem "
          "os acompanha entram pela porta à direita da C, sem fila.", tam=10.5,
          cor=TINTA)
    caixa(s, 0.62, 7.1, 12.1, 0.24,
          "O Hall 2, o apron e o pátio do Ring 3 estão em escala e na posição relativa "
          "medida em planta. O trecho entre o portão e o apron é esquemático: indica o "
          "lado por onde se anda, não a distância.", tam=8, cor=CINZA)
    return s


# --------------------------------------------------------------------------
# Slide 3 — o Hall 2
# --------------------------------------------------------------------------
def slide_hall(prs, planta, dec, mesas):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s)
    rotulo_secao(s, 0.62, 0.36, "Dentro do salão")
    caixa(s, 0.62, 0.6, 6.0, 0.4, "O Hall 2", tam=25, negrito=True)
    caixa(s, 3.1, 0.66, 6.2, 0.6,
          "As mesas não têm número: cada uma traz as seções que votam nela, que é o "
          "que você tem na mão.", tam=10.5, cor=CINZA)

    LARG, ALT = planta["salao"]["largura"], planta["salao"]["altura"]
    PROFM, LARGM = planta["modulo"]["prof"], planta["modulo"]["larg"]
    LX, LY, LH = 0.62, 1.15, 5.52
    e = LH / ALT
    X = lambda x: LX + x * e
    Y = lambda y: LY + (ALT - y) * e

    retangulo(s, X(0), Y(ALT), LARG * e, ALT * e, cor=BRANCO, borda=TINTA, larg=1.5)

    for par, k in (("oeste", "A"), ("norte", "B"), ("leste", "C")):
        banda = SF.BANDA_PAREDE[par]
        ys = [m["y"] for m in mesas if m["parede"] == par]
        xs = [m["x"] for m in mesas if m["parede"] == par]
        if par == "oeste":
            r = (PROFM, min(ys) - 1.4, banda, max(ys) + 1.4)
        elif par == "leste":
            r = (47.3 - banda, min(ys) - 1.4, 47.3 - PROFM, max(ys) + 1.4)
        else:
            r = (min(xs) - 1.4, ALT - banda, max(xs) + 1.4, ALT - PROFM)
        retangulo(s, X(r[0]), Y(r[3]), (r[2] - r[0]) * e, (r[3] - r[1]) * e,
                  cor=ZONA[k], alpha=0.88)

    for k, av in SF.AVENIDAS.items():
        # a avenida é a faixa entre os dois trilhos; desenhada por trechos retos
        ti, te = av["trilho_interno"], av["trilho_externo"]
        for (a1, a2), (b1, b2) in zip(zip(ti, ti[1:]), zip(te, te[1:])):
            xa = min(a1[0], a2[0], b1[0], b2[0]); xb = max(a1[0], a2[0], b1[0], b2[0])
            ya = min(a1[1], a2[1], b1[1], b2[1]); yb = max(a1[1], a2[1], b1[1], b2[1])
            retangulo(s, X(xa), Y(yb), (xb - xa) * e, (yb - ya) * e, cor=ZONA[k],
                      alpha=0.76)
        exi = (ti[0][0] + te[0][0]) / 2
        linha(s, X(exi), Y(1.0), X(exi), Y(7.5), cor=ZONA[k], larg=2, seta=True)
    xa, xb = SF.X_AVENIDA_NORTE
    retangulo(s, X(xa), Y(SF.Y_BANDA_NORTE), (xb - xa) * e, SF.AVENIDA_NORTE * e,
              cor=ZONA["B"], alpha=0.74)

    for m in mesas:
        k = m["entrada"]
        dxm, dym = m["dx"], m["dy"]
        xs = [m["x"], m["x"] + dxm * PROFM]; ys = [m["y"], m["y"] + dym * PROFM]
        if m["parede"] == "norte":
            x1, x2 = m["x"] - LARGM / 2, m["x"] + LARGM / 2
            y1, y2 = min(ys), max(ys)
        else:
            x1, x2 = min(xs), max(xs)
            y1, y2 = m["y"] - LARGM / 2, m["y"] + LARGM / 2
        retangulo(s, X(x1), Y(y2), (x2 - x1) * e, (y2 - y1) * e, cor=ZONA[k],
                  borda=TINTA, larg=0.5)
        rot = " ".join(MP.sec(x) for x in m["secoes"])
        cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
        if m["parede"] == "norte":
            caixa(s, X(cx) - (y2 - y1) * e / 2, Y(cy) - (x2 - x1) * e / 2,
                  (y2 - y1) * e, (x2 - x1) * e, rot, tam=6.5, negrito=True,
                  cor=TINTA_Z[k], alinha=PP_ALIGN.CENTER, ancora=MSO_ANCHOR.MIDDLE,
                  giro=-90)
        else:
            caixa(s, X(x1), Y(y2), (x2 - x1) * e, (y2 - y1) * e, rot, tam=6.5,
                  negrito=True, cor=TINTA_Z[k], alinha=PP_ALIGN.CENTER,
                  ancora=MSO_ANCHOR.MIDDLE)

    for p in planta["portas"]:
        sin = dec["sinalizacao_portas"].get(p["id"], {})
        papel = sin.get("papel")
        if papel not in ("entrada", "saida", "preferencial"):
            continue
        cor = ZONA[sin["entrada"]] if papel == "entrada" else (
            PREF if papel == "preferencial" else GRAFITE)
        linha(s, X(p["x1"]), Y(0), X(p["x2"]), Y(0), cor=cor, larg=6)
        cxp = (p["x1"] + p["x2"]) / 2
        if papel == "entrada":
            caixa(s, X(cxp) - 0.3, Y(0) + 0.05, 0.6, 0.26, sin["entrada"], tam=15,
                  negrito=True, cor=TXT_Z[sin["entrada"]], alinha=PP_ALIGN.CENTER)
            caixa(s, X(cxp) - 0.4, Y(0) + 0.28, 0.8, 0.2, "entrada", tam=7,
                  cor=CINZA, alinha=PP_ALIGN.CENTER)
        else:
            rotu = "preferencial" if papel == "preferencial" else "saída"
            corr = PREF if papel == "preferencial" else CINZA
            caixa(s, X(cxp) - 0.45, Y(0) + 0.06, 0.9, 0.2, rotu, tam=7.5,
                  negrito=True, cor=corr, alinha=PP_ALIGN.CENTER)

    # legenda, à direita
    lx = X(LARG) + 0.35
    caixa(s, lx, 1.15, 5.3, 0.24, "Como ler", tam=11, negrito=True, esp=1.2,
          maiuscula=True)
    ly = 1.55
    itens = [(ZONA["A"], "Porta A → parede oeste"),
             (ZONA["B"], "Porta B → parede norte"),
             (ZONA["C"], "Porta C → parede leste"),
             (PREF, "Entrada preferencial, sem fila"),
             (GRAFITE, "Saída — pela mesma faixa da parede")]
    for cor, txt in itens:
        retangulo(s, lx, ly + 0.03, 0.17, 0.17, cor=cor, raio=0.3)
        caixa(s, lx + 0.32, ly, 4.9, 0.4, txt, tam=11.5, cor=TINTA)
        ly += 0.40
    ly += 0.15
    caixa(s, lx, ly, 5.3, 2.6,
          "Da porta, a avenida da sua cor leva até a parede. Da avenida da parede, um "
          "passo até a sua seção.\n\n"
          "Quem votou volta pela mesma faixa da parede e sai pela fachada sul — "
          "ninguém atravessa a avenida de quem está entrando.\n\n"
          "A faixa mais escura na frente da parede norte é a pequena avenida de "
          "1,80 m: é ali que a fila da porta B escolhe o lado, andando.",
          tam=10, cor=CINZA)
    caixa(s, 0.62, 7.24, 12.1, 0.22,
          "Planta em escala · fachada sul embaixo, como quem chega do Ring 3 · "
          "todas as formas são editáveis", tam=8, cor=CINZA)
    return s


# --------------------------------------------------------------------------
# Slide 4 — sua seção, sua porta
# --------------------------------------------------------------------------
def slide_secoes(prs, por_porta, dados):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fundo(s)
    rotulo_secao(s, 0.62, 0.42, "Sua seção, sua porta")
    caixa(s, 0.62, 0.68, 8.0, 0.5, "Ache o seu número", tam=27, negrito=True)
    caixa(s, 0.62, 1.22, 9.5, 0.3,
          f"As {dados['secoes']} seções da Irlanda, repartidas pelas três portas. "
          "Anote a letra: você vai procurá-la três vezes no caminho.", tam=11,
          cor=CINZA)
    lx, lw = 0.62, 4.0
    for k in "ABC":
        linhas = -(-len(por_porta[k]) // 4)
        retangulo(s, lx, 1.75, lw, 1.10 + linhas * 0.42 + 0.22, cor=BRANCO, borda=REGUA)
        retangulo(s, lx, 1.75, lw, 0.07, cor=ZONA[k])
        retangulo(s, lx + 0.25, 2.05, 0.52, 0.52, cor=ZONA[k], raio=0.16)
        caixa(s, lx + 0.25, 2.16, 0.52, 0.32, k, tam=22, negrito=True,
              cor=TINTA_Z[k], alinha=PP_ALIGN.CENTER)
        caixa(s, lx + 0.9, 2.1, 3.0, 0.24, f"porta {MP.PORTA[k]}", tam=13,
              negrito=True)
        caixa(s, lx + 0.9, 2.34, 3.0, 0.24,
              f"parede {MP.PAREDE[k]} · {len(por_porta[k])} seções", tam=10,
              cor=CINZA)
        cx, cy, col = lx + 0.28, 2.85, 0
        for n in por_porta[k]:
            retangulo(s, cx + col * 0.86, cy, 0.76, 0.33, cor=ZONA[k], raio=0.18)
            caixa(s, cx + col * 0.86, cy + 0.055, 0.76, 0.24, MP.sec(n), tam=13,
                  negrito=True, cor=TINTA_Z[k], alinha=PP_ALIGN.CENTER)
            col += 1
            if col == 4:
                col, cy = 0, cy + 0.42
        lx += 4.25
    caixa(s, 0.62, 5.5, 12.1, 0.5,
          "Consulte sempre o seu local e a sua seção no site do TSE antes de sair de "
          "casa. · Gerado por scripts/mapa_publico_pptx.py a partir da planta medida "
          "do Hall 2 e da agregação oficial de seções.", tam=8.5, cor=CINZA)
    return s


# --------------------------------------------------------------------------
def monta():
    planta, dec, cen = SF.carrega()
    MP.APRON = planta["decisoes"]["ring3"]["apron"]
    mesas = SF.monta_mesas(planta, dec, cen)
    por_porta = {k: sorted({s for m in mesas if m["entrada"] == k for s in m["secoes"]})
                 for k in "ABC"}
    total = sum(len(v) for v in por_porta.values())
    assert total == 51, f"as seções não fecham em 51: {total}"
    dados = {
        "secoes": total, "urnas": len(mesas),
        "aptos": f"{dec['comparecimento']['aptos']:,}".replace(",", "."),
        "passos": [
            ("Portão B, na Merrion Road",
             "É por ali que o eleitor entra. Antes do portão há uma mesa de consulta: "
             "quem não sabe o número da sua seção descobre ali, e anota a letra."),
            ("Pela lateral leste do Hall 2",
             "O caminho contorna o salão pelo lado leste e desemboca no apron, entre "
             "o Hall 2 e o pátio."),
            ("Ring 3, pelo canto nordeste",
             "O pátio de fila fica ao ar livre, ao sul do salão. O corredor de chegada "
             "distribui para as três zonas: C primeiro, B depois, A no fim."),
            ("A sua zona, a sua fila",
             "Cada zona é a fila de uma porta, e as três têm o mesmo tamanho. Entra-se "
             "pela boca no fundo e sai-se pela abertura da frente."),
            ("A porta A, B ou C",
             "Atravessado o apron, cada zona entra pela sua porta. Da porta até a mesa "
             "é sempre a mesma letra."),
        ],
    }

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W_SLIDE), Inches(H_SLIDE)
    capa(prs, dados)
    slide_caminho(prs, planta, dec, dados)
    slide_hall(prs, planta, dec, mesas)
    slide_secoes(prs, por_porta, dados)
    return prs


def main():
    prs = monta()
    n = sum(len(s.shapes) for s in prs.slides)
    print(f"{len(prs.slides.__iter__.__self__._sldIdLst)} slides · {n} formas nativas")
    if "--grava" not in sys.argv[1:]:
        print("(sem --grava: nada foi escrito)")
        return 0
    prs.save(SAIDA)
    print("gravado", os.path.relpath(SAIDA, MP.RAIZ), f"({SAIDA.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
